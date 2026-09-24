"""Parse object inputs without publishing or depending on the private source workbook."""

import csv
import hashlib
import io
import zipfile
from copy import deepcopy
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from demo.snapshots import PARAMETER_LABELS


SHEETS = {
    "warehouse": "Склад",
    "airport": "Аэропорт",
    "hospital": "Медучреждение",
}
CSV_COLUMNS = ("object_slug", "label", "unit", "value", "min", "max", "source")
MAX_UPLOAD_BYTES = 1_000_000
MAX_UNCOMPRESSED_BYTES = 10_000_000
MAX_ROWS = 200
UNITS = {
    "-", "%", "SKU", "°C", "дБА", "ед./сут", "заявок/сут",
    "дн.", "кВт", "кг", "кг/сут", "коек", "конт./сут", "лет", "м",
    "м/п", "м²", "млн пасс./год", "млн руб.", "мм", "мм/2м",
    "мин", "наименований", "операций", "пасс./сут", "пасс./ч",
    "поддон/сут", "порций/сут", "посещений/сут", "проб/сут",
    "раз/сут", "рейсов/сут", "рейсов/ч", "руб./мес.",
    "смен", "смен/сут", "строк/сут", "строк/ч·чел", "ч", "чел.",
    "шт.", "шт./сут",
}


class ProfileValidationError(ValueError):
    pass


def _number(raw, label):
    try:
        number = Decimal(str(raw).strip().replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ProfileValidationError(f"{label}: требуется число") from None
    if not number.is_finite():
        raise ProfileValidationError(f"{label}: требуется конечное число")
    return number


def _json_number(number):
    return int(number) if number == number.to_integral_value() else float(number)


def normalize_value(raw, kind, minimum=None, maximum=None):
    if raw is None or str(raw).strip() == "":
        return None
    if kind == "boolean":
        if raw in (True, "true", "True", "Да"):
            return True
        if raw in (False, "false", "False", "Нет"):
            return False
        raise ProfileValidationError("Для логического поля выберите Да или Нет")
    if kind == "number":
        number = _number(raw, "Значение")
        if minimum is not None and number < Decimal(str(minimum)):
            raise ProfileValidationError(f"Значение меньше минимума {minimum}")
        if maximum is not None and number > Decimal(str(maximum)):
            raise ProfileValidationError(f"Значение больше максимума {maximum}")
        return _json_number(number)
    if isinstance(raw, (str, int, float)) and len(str(raw)) <= 200:
        value = str(raw).strip()
        if isinstance(minimum, str) and isinstance(maximum, str):
            if value not in {minimum, maximum}:
                raise ProfileValidationError("Значение вне допустимых текстовых вариантов")
        return value
    raise ProfileValidationError("Текстовое значение слишком длинное или имеет неверный тип")


def _field(row, row_number, *, formula=None, sheet=None):
    label, unit, raw, minimum, maximum, source = row[:6]
    if not isinstance(label, str) or not label.strip() or len(label) > 200:
        raise ProfileValidationError("Нет корректного названия параметра")
    if not isinstance(unit, str) or unit.strip() not in UNITS:
        raise ProfileValidationError(f"Неизвестная единица: {unit}")
    minimum = None if minimum is None or str(minimum).strip() in ("", "-") else minimum
    maximum = None if maximum is None or str(maximum).strip() in ("", "-") else maximum
    text_bounds = False
    if minimum is not None or maximum is not None:
        try:
            minimum = _json_number(_number(minimum, "Минимум")) if minimum is not None else None
            maximum = _json_number(_number(maximum, "Максимум")) if maximum is not None else None
        except ProfileValidationError:
            if unit.strip() != "-" or not isinstance(minimum, str) or not isinstance(maximum, str):
                raise
            minimum, maximum = minimum.strip(), maximum.strip()
            text_bounds = True
    if minimum is not None and maximum is not None and not text_bounds and minimum > maximum:
        raise ProfileValidationError("Минимум больше максимума")
    kind = "text" if text_bounds else "number" if minimum is not None or maximum is not None or isinstance(raw, (int, float)) else "text"
    value = normalize_value(raw, kind, minimum, maximum)
    if source is not None and len(str(source)) > 500:
        raise ProfileValidationError("Описание источника слишком длинное")
    return {
        "key": f"row-{row_number}",
        "label": label.strip(),
        "unit": unit.strip(),
        "kind": kind,
        "raw_value": formula if formula is not None else raw,
        "value": value,
        "min": minimum,
        "max": maximum,
        "status": "missing" if value is None else "assumption",
        "source": str(source).strip() if source is not None else "Пользовательский импорт",
        "source_sheet": sheet,
        "source_row": row_number,
        "formula": formula,
        "override": False,
    }


def profile_from_demo(snapshot):
    fields = []
    for key, data in snapshot["parameters"].items():
        value = data.get("value")
        fields.append({
            "key": key,
            "label": PARAMETER_LABELS.get(key, key),
            "unit": data.get("unit", "-"),
            "kind": "boolean" if key == "access_permission" else "number" if isinstance(value, (int, float)) and not isinstance(value, bool) else "text",
            "raw_value": value,
            "value": value,
            "min": None,
            "max": None,
            "status": data.get("status", "missing"),
            "source": data.get("source", "Синтетический пример"),
            "source_sheet": None,
            "source_row": None,
            "formula": None,
            "override": False,
        })
    return {"version": 1, "object_slug": snapshot["slug"], "source_type": "demo", "fields": fields}


def profile_for_revision(revision):
    snapshot = revision.scenario_snapshot
    return deepcopy(snapshot.get("input_profile") or profile_from_demo(snapshot))


def _parse_rows(rows, object_slug, source_type, digest, sheet=None, formulas=None):
    fields, errors = [], []
    if len(rows) > MAX_ROWS:
        raise ProfileValidationError("Слишком много строк в файле")
    for row_number, row in rows:
        if not row or row[0] is None or row[0] == "":
            continue
        if source_type == "xlsx" and len(row) >= 3 and row[1] is None and row[2] is None:
            continue  # Section headings in the supplied workbook.
        try:
            formula = formulas.get(row_number) if formulas else None
            if formula is not None and row[2] is None:
                raise ProfileValidationError("У формулы нет сохранённого значения")
            fields.append(_field(row, row_number, formula=formula, sheet=sheet))
        except ProfileValidationError as exc:
            errors.append(f"Строка {row_number}: {exc}")
    if not fields and not errors:
        raise ProfileValidationError("В файле нет параметров")
    return {
        "version": 1,
        "object_slug": object_slug,
        "source_type": source_type,
        "source_sha256": digest,
        "fields": fields,
        "errors": errors,
    }


def parse_upload(upload, object_slug):
    if upload.size > MAX_UPLOAD_BYTES:
        raise ProfileValidationError("Файл превышает предел 1 МБ")
    name = upload.name.lower()
    if not name.endswith((".xlsx", ".csv")):
        raise ProfileValidationError("Допустимы только .xlsx и .csv")
    content = upload.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise ProfileValidationError("Файл превышает предел 1 МБ")
    digest = hashlib.sha256(content).hexdigest()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text), delimiter=";")
            if tuple(reader.fieldnames or ()) != CSV_COLUMNS:
                raise ProfileValidationError("CSV должен иметь колонки object_slug;label;unit;value;min;max;source")
            rows = []
            for number, record in enumerate(reader, 2):
                if None in record or any(record[column] is None for column in CSV_COLUMNS):
                    raise ProfileValidationError(f"Строка {number}: неверное число колонок")
                if record["object_slug"] != object_slug:
                    raise ProfileValidationError(f"Строка {number}: тип объекта не соответствует проекту")
                rows.append((number, tuple(record.get(column) for column in CSV_COLUMNS[1:])))
        except UnicodeDecodeError:
            raise ProfileValidationError("CSV должен быть в кодировке UTF-8") from None
        return _parse_rows(rows, object_slug, "csv", digest)
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ProfileValidationError("Повреждённый XLSX")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > 100 or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES:
            raise ProfileValidationError("XLSX слишком велик после распаковки")
    try:
        values_book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formulas_book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheet = SHEETS[object_slug]
        if sheet not in values_book or sheet not in formulas_book:
            raise ProfileValidationError(f"В XLSX нет листа «{sheet}»")
        last_row = values_book[sheet].max_row
        if last_row > MAX_ROWS + 2:
            raise ProfileValidationError("Слишком много строк в файле")
        values = list(values_book[sheet].iter_rows(min_row=3, max_row=last_row, max_col=6, values_only=True))
        formula_rows = formulas_book[sheet].iter_rows(min_row=3, max_row=last_row, max_col=6)
        formulas = {}
        for number, row in enumerate(formula_rows, 3):
            if row[2].data_type == "f":
                formulas[number] = row[2].value
        rows = [(number, row) for number, row in enumerate(values, 3)]
        return _parse_rows(rows, object_slug, "xlsx", digest, sheet=sheet, formulas=formulas)
    except Exception as exc:  # Untrusted XML/ZIP can fail in several parser layers.
        if isinstance(exc, ProfileValidationError):
            raise
        raise ProfileValidationError("XLSX не удалось прочитать") from None
    finally:
        if "values_book" in locals():
            values_book.close()
        if "formulas_book" in locals():
            formulas_book.close()
