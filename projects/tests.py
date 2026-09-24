from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, TestCase
from openpyxl import Workbook

from projects.models import Project
from projects.input_profiles import profile_for_revision


class InputProfileTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", verbosity=0)
        cls.owner = get_user_model().objects.create_user(
            username="profile-owner", password="temporary-test-passphrase"
        )
        cls.other = get_user_model().objects.create_user(
            username="profile-other", password="temporary-test-passphrase"
        )

    def setUp(self):
        self.client.force_login(self.owner)

    def create_project(self, slug):
        response = self.client.post(
            f"/projects/new/{slug}/", {"name": f"Профиль {slug}"}
        )
        self.assertEqual(response.status_code, 302)
        return Project.objects.get(id=response["Location"].split("/")[2])

    def workbook(self):
        book = Workbook()
        for number, (sheet, unit) in enumerate((
            ("Склад", "рейсов/сут"),
            ("Аэропорт", "конт./сут"),
            ("Медучреждение", "заявок/сут"),
        )):
            page = book.active if number == 0 else book.create_sheet()
            page.title = sheet
            page.append(["Синтетический тестовый объект"])
            page.append(["Параметр", "Ед. изм.", "Значение", "Мин", "Макс", "Источник"])
            page.append(["Раздел"])
            page.append(["Синтетическая нагрузка", unit, 10, 0, 20, "Тестовый набор"])
        output = BytesIO()
        book.save(output)
        return output.getvalue()

    def upload(self, project, content, name="synthetic.xlsx"):
        return self.client.post(
            f"/projects/{project.id}/import/",
            {"action": "preview", "file": SimpleUploadedFile(name, content)},
        )

    def form_values(self, project):
        profile = profile_for_revision(project.revisions.first())
        values = {
            f"field_{i}": (
                "true" if field["value"] is True else "false" if field["value"] is False
                else field["value"] if field["value"] is not None else ""
            )
            for i, field in enumerate(profile["fields"])
        }
        values["base_revision"] = str(project.revisions.first().number)
        return values

    def test_owner_can_open_inputs_for_each_object(self):
        for slug in ("warehouse", "airport", "hospital"):
            with self.subTest(slug=slug):
                project = self.create_project(slug)
                page = self.client.get(f"/projects/{project.id}/inputs/")
                self.assertEqual(page.status_code, 200)
                self.assertContains(page, project.get_object_slug_display())

    def test_edit_creates_revision_and_preserves_original(self):
        for slug in ("warehouse", "airport", "hospital"):
            with self.subTest(slug=slug):
                project = self.create_project(slug)
                original = project.revisions.get(number=1)
                data = self.form_values(project)
                data["field_0"] = "77"
                data["base_revision"] = "1"
                response = self.client.post(f"/projects/{project.id}/inputs/", data)
                self.assertEqual(response.status_code, 302)
                self.assertEqual(project.revisions.count(), 2)
                updated = profile_for_revision(project.revisions.get(number=2))
                self.assertEqual(updated["fields"][0]["value"], 77)
                self.assertTrue(updated["fields"][0]["override"])
                original.refresh_from_db()
                self.assertNotEqual(profile_for_revision(original)["fields"][0]["value"], 77)
                self.assertContains(self.client.get(f"/projects/{project.id}/?revision=1"), "Ревизия 1")

    def test_xlsx_preview_and_commit_for_three_sheets(self):
        content = self.workbook()
        for slug in ("warehouse", "airport", "hospital"):
            with self.subTest(slug=slug):
                project = self.create_project(slug)
                preview = self.upload(project, content)
                self.assertEqual(preview.status_code, 200)
                self.assertContains(preview, "Синтетическая нагрузка")
                self.assertEqual(project.revisions.count(), 1)
                saved = self.client.post(
                    f"/projects/{project.id}/import/", {"action": "commit"}
                )
                self.assertEqual(saved.status_code, 302)
                profile = profile_for_revision(project.revisions.get(number=2))
                self.assertEqual(profile["object_slug"], slug)
                self.assertEqual(profile["fields"][0]["value"], 10)
                self.assertEqual(profile["fields"][0]["source_row"], 4)
                self.assertEqual(len(profile["source_sha256"]), 64)

    def test_csv_wrong_object_and_out_of_range_are_rejected(self):
        project = self.create_project("warehouse")
        heading = "object_slug;label;unit;value;min;max;source\n"
        wrong = (heading + "airport;Спрос;рейсов/сут;10;0;20;Тест\n").encode()
        self.assertEqual(self.upload(project, wrong, "wrong.csv").status_code, 400)
        high = (heading + "warehouse;Спрос;рейсов/сут;21;0;20;Тест\n").encode()
        response = self.upload(project, high, "high.csv")
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "больше максимума", status_code=400)
        self.assertEqual(project.revisions.count(), 1)
        self.assertEqual(
            self.client.post(f"/projects/{project.id}/import/", {"action": "commit"}).status_code,
            400,
        )
        valid = (heading + "warehouse;Спрос;рейсов/сут;0;0;20;Тест\n").encode()
        self.assertEqual(self.upload(project, valid, "valid.csv").status_code, 200)
        self.assertEqual(self.client.post(
            f"/projects/{project.id}/import/", {"action": "commit"}
        ).status_code, 302)
        self.assertEqual(profile_for_revision(project.revisions.get(number=2))["fields"][0]["value"], 0)
        invalid_edit = self.form_values(project)
        invalid_edit["field_0"] = "21"
        self.assertEqual(
            self.client.post(f"/projects/{project.id}/inputs/", invalid_edit).status_code,
            400,
        )
        self.assertEqual(project.revisions.count(), 2)

    def test_invalid_upload_does_not_save_revision(self):
        project = self.create_project("hospital")
        for name, content in (
            ("bad.txt", b"text"),
            ("bad.xlsx", b"not a zip"),
            ("large.csv", b"x" * 1_000_001),
            ("bad.csv", b"object_slug;label;unit;value;min;max;source\nhospital;X;unknown;1;0;2;Test\n"),
        ):
            with self.subTest(name=name):
                self.assertEqual(self.upload(project, content, name).status_code, 400)
                self.assertEqual(project.revisions.count(), 1)

    def test_textual_limits_formula_without_cache_and_stale_revision(self):
        project = self.create_project("warehouse")
        csv_file = (
            "object_slug;label;unit;value;min;max;source\n"
            "warehouse;Доступ;-;Да;Да;Да;Синтетический тест\n"
        ).encode()
        self.assertEqual(self.upload(project, csv_file, "text.csv").status_code, 200)
        self.assertEqual(self.client.post(
            f"/projects/{project.id}/import/", {"action": "commit"}
        ).status_code, 302)
        values = self.form_values(project)
        values["field_0"] = "Нет"
        response = self.client.post(f"/projects/{project.id}/inputs/", values)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(project.revisions.count(), 2)

        values["field_0"] = "Да"
        self.assertEqual(self.client.post(f"/projects/{project.id}/inputs/", values).status_code, 302)
        self.assertEqual(project.revisions.count(), 2)

        book = Workbook()
        book.active.title = "Склад"
        book.active.append(["Тест"])
        book.active.append(["Параметр", "Единица", "Значение", "Мин", "Макс", "Источник"])
        book.active.append(["Нагрузка", "рейсов/сут", "=5+5", 0, 20, "Синтетический тест"])
        output = BytesIO()
        book.save(output)
        preview = self.upload(project, output.getvalue())
        self.assertEqual(preview.status_code, 400)
        self.assertContains(preview, "нет сохранённого значения", status_code=400)

        project2 = self.create_project("airport")
        first = self.form_values(project2)
        first["field_0"] = "55"
        self.assertEqual(self.client.post(f"/projects/{project2.id}/inputs/", first).status_code, 302)
        old = dict(first, field_0="56")
        stale = self.client.post(f"/projects/{project2.id}/inputs/", old)
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(project2.revisions.count(), 2)

    def test_ownership_guest_and_csrf(self):
        project = self.create_project("warehouse")
        path = f"/projects/{project.id}/import/"
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(path).status_code, 404)
        self.assertEqual(self.client.post(path, {"action": "commit"}).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(path).status_code, 302)
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.owner)
        self.assertEqual(csrf_client.post(path, {"action": "commit"}).status_code, 403)

    def test_preview_escapes_uploaded_text(self):
        project = self.create_project("hospital")
        content = (
            "object_slug;label;unit;value;min;max;source\n"
            "hospital;<script>alert(1)</script>;шт.;1;0;2;Синтетический тест\n"
        ).encode()
        page = self.upload(project, content, "escape.csv")
        self.assertEqual(page.status_code, 200)
        self.assertNotIn(b"<script>", page.content)
        self.assertIn(b"&lt;script&gt;", page.content)
